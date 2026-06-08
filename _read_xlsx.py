import openpyxl

wb = openpyxl.load_workbook(r"E:\proj\web\42ium.xyz_1780565437970.xlsx")
print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== {name} ({ws.max_row} rows x {ws.max_column} cols) ===")
    for row in ws.iter_rows(min_row=1, values_only=True):
        print(" | ".join(str(c) if c is not None else "" for c in row))
